import copy
import json
from pathlib import Path

import httpx
import pytest
from docx import Document
from pypdf import PdfReader

from jobpilot.documents import choose, prepare, resume_sections


def candidate():
    return {
        "name": "Alex Example",
        "email": "alex@example.com",
        "location": "Columbus, OH",
        "experience": [
            {
                "company": "Example Systems",
                "title": "Security Engineer",
                "start": "January 2024",
                "end": "Present",
                "location": "Columbus, OH",
                "facts": [
                    {
                        "id": "access",
                        "text": "Reviewed AWS IAM policies and automated account access reviews.",
                    },
                    {
                        "id": "meetings",
                        "text": "Organized weekly engineering status meetings and maintained notes.",
                    },
                    {
                        "id": "audit",
                        "text": "Maintained audit evidence with engineering and business colleagues.",
                    },
                ],
            },
            {
                "company": "Example Labs",
                "title": "Security Analyst",
                "start": "June 2022",
                "end": "December 2023",
                "location": "Dayton, OH",
                "facts": [
                    {
                        "id": "alerts",
                        "text": "Reviewed security alerts and documented investigation findings.",
                    },
                    {
                        "id": "inventory",
                        "text": "Maintained the equipment inventory and routed support tickets.",
                    },
                ],
            },
        ],
        "skills": ["AWS", "IAM", "Linux", "Office administration"],
        "certifications": ["Example Security Certificate"],
        "education": [
            {"institution": "Example College", "degree": "Bachelor of Science", "end": "2022"}
        ],
        "interests": ["Chess", "Cooking & baking"],
    }


JOB = {
    "id": 18,
    "title": "Cloud Security Engineer",
    "company": "Future Company",
    "description": "AWS IAM Terraform ISO27001 audit evidence. Deliver a roadmap in 90 days.",
}
SETTINGS = {"ollama_url": "http://127.0.0.1:11436", "model": "qwen3.5:9b"}


def local_response(monkeypatch, plan):
    requests = []
    original_client = httpx.Client

    def respond(request):
        if request.url.path == "/api/tags":
            return httpx.Response(200, json={"models": [{"name": "qwen3.5:9b"}]})
        payload = json.loads(request.content)
        requests.append(payload)
        # Leave the independent cover-letter draft on its existing template fallback.
        content = {"paragraphs": []} if "paragraphs" in payload["format"]["properties"] else plan
        return httpx.Response(200, json={"message": {"content": json.dumps(content)}})

    monkeypatch.setattr(
        "jobpilot.documents.httpx.Client",
        lambda **kwargs: original_client(transport=httpx.MockTransport(respond), **kwargs),
    )
    return requests


def test_local_resume_can_select_rewrite_and_replace_skills_in_both_exports(monkeypatch, tmp_path):
    profile = candidate()
    original = copy.deepcopy(profile)
    access = "Reviewed AWS IAM policies using Terraform to automate account access reviews."
    audit = "Maintained ISO27001 audit evidence with engineering colleagues."
    requests = local_response(
        monkeypatch,
        {
            "fact_ids": ["audit", "access", "alerts"],
            "bullet_edits": [
                {"fact_id": "audit", "text": audit},
                {"fact_id": "access", "text": access},
            ],
            "skills": ["Terraform", "AWS", "ISO 27001"],
            "gaps": [],
        },
    )
    result = prepare(JOB, profile, tmp_path, SETTINGS)
    assert result["generation_mode"] == "ollama"
    texts = [
        " ".join(
            " ".join(page.extract_text().split()) for page in PdfReader(result["resume_pdf"]).pages
        ),
        " ".join(paragraph.text for paragraph in Document(result["resume_docx"]).paragraphs),
    ]
    for text in texts:
        assert access in text and audit in text
        assert text.index(audit) < text.index(access)
        assert "weekly engineering status meetings" not in text
        assert "equipment inventory" not in text
        assert "Skills: Terraform, AWS, ISO 27001" in text
        assert "Linux" not in text and "Office administration" not in text
        for fixed in (
            "Example Systems",
            "Security Engineer",
            "January 2024",
            "Present",
            "Example Labs",
            "Security Analyst",
            "June 2022",
            "December 2023",
            "Example College",
            "Bachelor of Science",
            "Example Security Certificate",
        ):
            assert fixed in text
        assert (
            text.index("SKILLS & CERTIFICATIONS")
            < text.index("Interests: Chess, Cooking & baking")
            < text.index("EDUCATION")
        )
    assert profile == original
    manifest = json.loads(Path(result["manifest"]).read_text())
    assert manifest["template_version"] == "1.2"
    assert manifest["fact_ids"] == ["audit", "access", "alerts"]
    assert manifest["career_claims"]["access"] == original["experience"][0]["facts"][0]["text"]
    assert manifest["rendered_bullets"]["access"] == access
    assert manifest["resume_edits"]["bullets"]["audit"] == audit
    assert manifest["resume_skills"] == ["Terraform", "AWS", "ISO 27001"]
    assert (
        len([request for request in requests if "fact_ids" in request["format"]["properties"]]) == 1
    )


def test_local_selection_keeps_one_fact_for_each_employer_without_restoring_all(monkeypatch):
    local_response(
        monkeypatch, {"fact_ids": ["access"], "bullet_edits": [], "skills": ["AWS"], "gaps": []}
    )
    selected, mode, _ = choose(JOB, candidate(), SETTINGS)
    assert mode == "ollama"
    assert selected == ["access", "alerts"]


@pytest.mark.parametrize("timing", ["within hours", "in a few days", "under an hour"])
def test_unsupported_timing_restores_only_affected_bullet(monkeypatch, tmp_path, timing):
    profile = candidate()
    original = profile["experience"][0]["facts"][0]["text"]
    audit = "Maintained ISO27001 audit evidence with engineering colleagues."
    local_response(
        monkeypatch,
        {
            "fact_ids": ["access", "audit", "alerts"],
            "bullet_edits": [
                {
                    "fact_id": "access",
                    "text": f"Automated AWS IAM account access reviews {timing}.",
                },
                {"fact_id": "audit", "text": audit},
            ],
            "skills": ["Terraform", "AWS"],
            "gaps": [],
        },
    )
    result = prepare(JOB, profile, tmp_path, SETTINGS)
    assert result["generation_mode"] == "ollama"
    manifest = json.loads(Path(result["manifest"]).read_text())
    assert manifest["fact_ids"] == ["access", "audit", "alerts"]
    assert manifest["rendered_bullets"]["access"] == original
    assert manifest["rendered_bullets"]["audit"] == audit
    assert manifest["resume_skills"] == ["Terraform", "AWS"]
    for text in (
        " ".join(page.extract_text() for page in PdfReader(result["resume_pdf"]).pages),
        " ".join(paragraph.text for paragraph in Document(result["resume_docx"]).paragraphs),
    ):
        assert timing not in text
        assert original in " ".join(text.split())
        assert audit in " ".join(text.split())


@pytest.mark.parametrize(
    ("fact_ids", "bullet_edits"),
    [
        (["access", "invented"], []),
        (
            ["access", "alerts"],
            [
                {
                    "fact_id": "access",
                    "text": "Reduced AWS security incidents by 987% using Terraform.",
                }
            ],
        ),
        # A number in the job is not evidence for a historical achievement.
        (
            ["access", "alerts"],
            [
                {
                    "fact_id": "access",
                    "text": "Completed AWS access reviews for 90 production systems.",
                }
            ],
        ),
    ],
)
def test_invalid_resume_claims_fall_back_to_all_original_facts(
    monkeypatch, tmp_path, fact_ids, bullet_edits
):
    local_response(
        monkeypatch,
        {"fact_ids": fact_ids, "bullet_edits": bullet_edits, "skills": ["Terraform"], "gaps": []},
    )
    result = prepare(JOB, candidate(), tmp_path, SETTINGS)
    assert result["generation_mode"] == "deterministic"
    manifest = json.loads(Path(result["manifest"]).read_text())
    assert set(manifest["fact_ids"]) == {"access", "meetings", "audit", "alerts", "inventory"}
    text = " ".join(
        " ".join(page.extract_text().split()) for page in PdfReader(result["resume_pdf"]).pages
    )
    for experience in candidate()["experience"]:
        for fact in experience["facts"]:
            assert fact["text"] in text
    assert "987%" not in text and "90 production systems" not in text


def test_deterministic_skills_drop_unrelated_saved_terms_and_keep_no_job_fallback():
    _, skills = resume_sections(candidate(), ["access", "alerts"], JOB)
    assert "Terraform" in skills and "AWS" in skills and "IAM" in skills
    assert "Linux" not in skills and "Office administration" not in skills
    _, no_job_skills = resume_sections(candidate(), ["access", "alerts"], {})
    assert no_job_skills == ["AWS", "IAM", "Linux", "Office administration"]
