import json
from pathlib import Path

import httpx
import pytest
from docx import Document
from pypdf import PdfReader

from jobpilot.documents import _draft_cover, choose, prepare
from jobpilot.keywords import extract_keywords


def substantial_profile():
    profile = json.loads((Path(__file__).parents[1] / "jobpilot/candidate.json").read_text())
    profile["experience"][0]["facts"] += [
        {
            "id": "cloud-config",
            "text": "Hardened AWS cloud configuration through infrastructure as code, including IAM policies, encryption controls and service-account access reviews.",
        },
        {
            "id": "vendor-reviews",
            "text": "Completed vendor security reviews and maintained audit evidence with engineering and business stakeholders.",
        },
    ]
    profile["experience"][1]["facts"] += [
        {
            "id": "automation",
            "text": "Automated vulnerability tracking and routed security findings to engineering teams for review and remediation.",
        },
        {
            "id": "access-workflow",
            "text": "Built a self-service access workflow with automatic revocation checks to support least-privilege administration.",
        },
    ]
    return profile


def test_selection_retains_substantial_experience_beyond_ten_facts():
    profile = substantial_profile()
    selected, _, _ = choose(
        {"description": "AWS IAM cloud security SOC 2"},
        profile,
        {"ollama_url": "https://disabled.example", "model": "cloud"},
    )
    assert len(selected) == 11


def test_resume_layout_and_keyword_metadata_preserve_facts_and_add_draft_tools(tmp_path):
    profile = substantial_profile()
    job = {
        "id": 12,
        "title": "Cloud Security Engineer",
        "company": "Example Company",
        "description": "AWS IAM cloud security Terraform Wiz Kubernetes SOC 2 risk assessment",
    }
    result = prepare(
        job, profile, tmp_path, {"ollama_url": "https://disabled.example", "model": "cloud"}
    )
    reader = PdfReader(result["resume_pdf"])
    assert len(reader.pages) == 1
    text = reader.pages[0].extract_text()
    assert "Northstar Systems" in text and "MapleWorks" in text
    assert "Terraform" in text and "Wiz" in text and "Kubernetes" in text
    assert "GPA 3.5" in text
    assert "\x7f" not in text
    for experience in profile["experience"]:
        for fact in experience["facts"]:
            assert fact["text"] in " ".join(text.split())
    assert {"Terraform", "Wiz", "Kubernetes"} <= set(result["keywords"]["suggested"])
    manifest = json.loads(Path(result["manifest"]).read_text())
    assert manifest["keywords"] == result["keywords"]
    assert len(manifest["fact_ids"]) == 11
    resume = Document(result["resume_docx"])
    assert any("\t" in paragraph.text for paragraph in resume.paragraphs)
    assert "January 2024" in " ".join(paragraph.text for paragraph in resume.paragraphs)


def test_cover_letter_is_prose_with_company_role_and_forward_looking_keyword_draft(tmp_path):
    profile = substantial_profile()
    job = {
        "title": "Cloud Security Engineer",
        "company": "Example Company",
        "description": "Terraform Wiz cloud security IAM SOC 2 incident response",
    }
    result = prepare(
        job, profile, tmp_path, {"ollama_url": "https://disabled.example", "model": "cloud"}
    )
    cover = Document(result["cover_docx"])
    text = "\n".join(paragraph.text for paragraph in cover.paragraphs)
    assert "Example Company" in text and "Cloud Security Engineer" in text
    assert "Terraform" in text and "Wiz" in text
    assert "I would" in text
    assert "Northstar Systems" in text and "MapleWorks" in text
    assert not any(paragraph.style.name == "List Bullet" for paragraph in cover.paragraphs)
    assert len(PdfReader(result["cover_pdf"]).pages) == 1


def test_missing_education_details_do_not_create_empty_gpa_or_date_labels(tmp_path):
    profile = substantial_profile()
    profile["education"] = [{"institution": "Example University", "degree": "Bachelor of Science"}]
    result = prepare(
        {"title": "Security Analyst", "company": "Example"},
        profile,
        tmp_path,
        {"ollama_url": "https://disabled.example", "model": "cloud"},
    )
    text = " ".join(page.extract_text() for page in PdfReader(result["resume_pdf"]).pages)
    assert "GPA" not in text
    assert "None" not in text


def cover_draft():
    return [
        "I am applying for the Cloud Security Engineer position at Example Company. "
        "My background in security engineering includes maintaining cloud controls and "
        "working with colleagues to make access and remediation processes usable. "
        "The role's emphasis on practical risk reduction fits the way I approach "
        "technical work: understand the system, agree on priorities, and document decisions.",
        "At Northstar Systems, I hardened AWS cloud configuration through infrastructure "
        "as code, including IAM policies and encryption controls. At MapleWorks, I "
        "automated vulnerability tracking and routed findings to engineering teams for "
        "review and remediation. These examples reflect my experience connecting "
        "security requirements with the people responsible for operating and maintaining systems.",
        "At Example Company, I would start by understanding existing cloud services and "
        "the team's most pressing security concerns. I would use Terraform and Wiz "
        "to support cloud configuration reviews, then work with engineers to prioritize "
        "findings and verify changes. Over the first 90 days, I would establish "
        "clear documentation and a practical review process for the team.",
    ]


def mock_cover_response(monkeypatch, paragraphs):
    requests = []
    original_client = httpx.Client

    def respond(request):
        requests.append(json.loads(request.content))
        return httpx.Response(
            200, json={"message": {"content": json.dumps({"paragraphs": paragraphs})}}
        )

    monkeypatch.setattr(
        "jobpilot.documents.httpx.Client",
        lambda **kwargs: original_client(transport=httpx.MockTransport(respond), **kwargs),
    )
    return requests


def test_local_cover_draft_accepts_prose_and_new_tools_with_job_numbers(monkeypatch):
    profile = substantial_profile()
    job = {
        "title": "Cloud Security Engineer",
        "company": "Example Company",
        "description": "Use Terraform and Wiz; establish the cloud security roadmap in 90 days.",
    }
    paragraphs = cover_draft()
    requests = mock_cover_response(monkeypatch, paragraphs)
    selected = [fact["id"] for experience in profile["experience"] for fact in experience["facts"]][
        ::-1
    ]
    result, mode = _draft_cover(
        job,
        profile,
        selected,
        extract_keywords(job, profile),
        {"ollama_url": "http://127.0.0.1:11436", "model": "qwen3.5:9b"},
    )
    assert mode == "ollama"
    assert result[4:7] == paragraphs
    assert "Terraform and Wiz" in result[6]
    assert requests[0]["think"] is False
    assert requests[0]["format"]["properties"]["paragraphs"]["maxItems"] == 3
    payload = json.loads(requests[0]["messages"][1]["content"])
    assert payload["experience"][0]["facts"][0] == profile["experience"][0]["facts"][-1]["text"]


@pytest.mark.parametrize(
    "claim",
    ["I reduced risk by 987%.", "I hold a CISSP credential.", "I have a decade of experience."],
)
def test_local_cover_draft_rejects_invented_past_metrics_and_credentials(monkeypatch, claim):
    profile = substantial_profile()
    job = {"title": "Cloud Security Engineer", "company": "Example Company"}
    paragraphs = cover_draft()
    paragraphs[1] += " " + claim
    # The number in the valid prospective paragraph is supplied by the job.
    job["description"] = "Establish the cloud security roadmap in 90 days."
    mock_cover_response(monkeypatch, paragraphs)
    result, mode = _draft_cover(
        job,
        profile,
        [fact["id"] for experience in profile["experience"] for fact in experience["facts"]],
        extract_keywords(job, profile),
        {"ollama_url": "http://127.0.0.1:11436", "model": "qwen3.5:9b"},
    )
    assert mode == "template"
    assert claim not in " ".join(result)


def test_local_cover_draft_rejects_a_truncated_paragraph(monkeypatch):
    profile = substantial_profile()
    job = {"title": "Cloud Security Engineer", "description": "Deliver a roadmap in 90 days."}
    paragraphs = cover_draft()
    paragraphs[2] = paragraphs[2][:-12]
    mock_cover_response(monkeypatch, paragraphs)
    result, mode = _draft_cover(
        job,
        profile,
        [fact["id"] for experience in profile["experience"] for fact in experience["facts"]],
        extract_keywords(job, profile),
        {"ollama_url": "http://127.0.0.1:11436", "model": "qwen3.5:9b"},
    )
    assert mode == "template"
    assert paragraphs[2] not in result
