import pytest
from docx import Document
from pypdf import PdfReader

from jobpilot.documents import write_docx, write_pdf


@pytest.mark.parametrize(
    ("interests", "expected"),
    [
        (
            ["  Guitar repair  ", "", None, "Woodwork & ceramics", "Classic cinema"],
            "Guitar repair, Woodwork & ceramics, Classic cinema",
        ),
        ("  Hiking & photography; Reading  ", "Hiking & photography; Reading"),
    ],
)
def test_resume_interests_follow_skills_and_certifications_in_both_formats(
    tmp_path, interests, expected
):
    profile = {
        "name": "Jordan Taylor",
        "interests": interests,
        "certifications": ["Example Certificate"],
        "education": [{"institution": "Example University", "degree": "Bachelor of Science"}],
    }
    pdf = tmp_path / "resume.pdf"
    docx = tmp_path / "resume.docx"
    write_pdf(pdf, profile, [], ["Python"])
    write_docx(docx, profile, [], ["Python"])
    texts = [
        " ".join(" ".join(page.extract_text().split()) for page in PdfReader(pdf).pages),
        " ".join(paragraph.text for paragraph in Document(docx).paragraphs),
    ]
    for text in texts:
        assert f"Interests: {expected}" in text
        assert text.index("SKILLS & CERTIFICATIONS") < text.index("Skills: Python")
        assert text.index("Certifications: Example Certificate") < text.index("Interests:")
        assert text.index("Interests:") < text.index("EDUCATION")
        assert "&amp;" not in text


def test_interests_are_retained_when_no_skills_or_certifications_are_available(tmp_path):
    profile = {"name": "Jordan Taylor", "interests": ["Chess", "Cooking & baking"]}
    pdf = tmp_path / "resume.pdf"
    docx = tmp_path / "resume.docx"
    write_pdf(pdf, profile, [])
    write_docx(docx, profile, [])
    texts = [
        " ".join(page.extract_text() for page in PdfReader(pdf).pages),
        " ".join(paragraph.text for paragraph in Document(docx).paragraphs),
    ]
    for text in texts:
        assert "SKILLS & CERTIFICATIONS" in text
        assert "Interests: Chess, Cooking & baking" in text


@pytest.mark.parametrize("interests", [None, "", "  ", [], [None, " "]])
def test_missing_interests_do_not_add_empty_or_invented_text(tmp_path, interests):
    profile = {"name": "Jordan Taylor", "interests": interests}
    pdf = tmp_path / "resume.pdf"
    docx = tmp_path / "resume.docx"
    write_pdf(pdf, profile, [], ["Python"])
    write_docx(docx, profile, [], ["Python"])
    texts = [
        " ".join(page.extract_text() for page in PdfReader(pdf).pages),
        " ".join(paragraph.text for paragraph in Document(docx).paragraphs),
    ]
    assert all("Interests:" not in text for text in texts)
